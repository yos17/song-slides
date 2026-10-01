"""Turn a plain-text song collection into a PowerPoint deck.

This is the single implementation shared by the command line (simple_generator.py)
and the browser app (web/, which runs this file through Pyodide).

Song file format: a line starting with "#" begins a song and holds its title.
Blank lines separate stanzas, and every stanza becomes one slide.
"""

from __future__ import annotations

import io
import math
from dataclasses import dataclass, field

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Inches, Pt

TITLE_SIZE_RANGE = (16, 72)
CONTENT_SIZE_RANGE = (12, 60)
MIN_FITTED_CONTENT_SIZE = 16  # long stanzas shrink down to this, then may overflow
MIN_FITTED_TITLE_SIZE = 20
COUNTER_SIZE = 24
TOC_ENTRY_SIZE = 20
TOC_MIN_ENTRY_SIZE = 12
TOC_MIN_SINGLE_LINE_SIZE = 14  # a long title shrinks alone, down to this, to stay on one line
TOC_LINE_GAP_EM = 6 / 20
TOC_COLUMNS = 2
TOC_SONGS_PER_COLUMN = 10
TOC_SONGS_PER_SLIDE = TOC_COLUMNS * TOC_SONGS_PER_COLUMN

BLACK = RGBColor(0, 0, 0)
COUNTER_BROWN = RGBColor(139, 69, 19)
LINK_BLUE = RGBColor(0, 0, 139)

# Geometry, tuned to sit below the red line of the church templates.
SIDE_MARGIN = Inches(0.5)
TEXT_INSET = Inches(0.2)
TITLE_TOP = Inches(0.6)
TITLE_HEIGHT = Inches(1.0)
COUNTER_WIDTH = Inches(1.0)
COUNTER_RIGHT = Inches(1.3)  # counter box starts this far from the right edge
CONTENT_TOP = Inches(1.4)
BOTTOM_MARGIN = Inches(0.8)  # clears the footer line of the church templates
TOC_TOP = Inches(1.6)
TOC_BOTTOM_MARGIN = Inches(1.2)
TOC_COLUMN_GAP = Inches(0.2)

# Rough text metrics used to estimate wrapping; Calibri-like fonts average
# about half an em per character.
CHAR_WIDTH_EM = 0.5
BOLD_CHAR_WIDTH_EM = 0.56
TOC_CHAR_WIDTH_EM = 0.42  # measured ~0.41 for Calibri (and metric-identical Carlito) titles
LINE_HEIGHT_EM = 1.2
STANZA_LINE_GAP_EM = 16 / 28  # 16pt after each line at the default 28pt


@dataclass
class Song:
    title: str
    stanzas: list[list[str]] = field(default_factory=list)


@dataclass
class Style:
    font_family: str = "Calibri"
    title_size: int = 32
    content_size: int = 28

    def __post_init__(self):
        self.font_family = (self.font_family or "").strip() or "Calibri"
        self.title_size = _clamp(int(self.title_size), *TITLE_SIZE_RANGE)
        self.content_size = _clamp(int(self.content_size), *CONTENT_SIZE_RANGE)


@dataclass
class Summary:
    songs: int = 0
    song_slides: int = 0
    toc_slides: int = 0
    warnings: list[str] = field(default_factory=list)

    @property
    def total(self) -> int:
        return self.song_slides + self.toc_slides

    def __str__(self) -> str:
        text = f"{self.total} slides from {self.songs} songs"
        if self.toc_slides:
            text += f" (including {self.toc_slides} table of contents slide{'s' if self.toc_slides > 1 else ''})"
        return text

    def to_dict(self) -> dict:
        return {
            "songs": self.songs,
            "song_slides": self.song_slides,
            "toc_slides": self.toc_slides,
            "total": self.total,
            "warnings": list(self.warnings),
        }


# --- parsing ---------------------------------------------------------------

def decode_text(data: bytes) -> str:
    """Decode an uploaded song file: UTF-8 (with or without BOM), else Windows-1252."""
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def parse_songs(text: str) -> list[Song]:
    """Parse "# Title" headed songs; blank lines split stanzas.

    Text before the first title and sections with an empty title are ignored.
    """
    text = text.lstrip("﻿").replace("\r\n", "\n").replace("\r", "\n")
    songs: list[Song] = []
    song: Song | None = None
    stanza: list[str] = []

    def end_stanza():
        if song is not None and stanza:
            song.stanzas.append(stanza.copy())
        stanza.clear()

    for raw in text.split("\n"):
        line = raw.rstrip()
        if line.startswith("#"):
            end_stanza()
            title = line.lstrip("#").strip()
            song = Song(title) if title else None
            if song:
                songs.append(song)
        elif line.strip():
            stanza.append(line)
        else:
            end_stanza()
    end_stanza()
    return songs


# --- text fitting ----------------------------------------------------------

def _wrapped_line_count(text: str, width_pt: float, size: float, char_em: float) -> int:
    chars_per_line = max(1, int(width_pt / (size * char_em)))
    return max(1, math.ceil(len(text) / chars_per_line))


def text_height(lines: list[str], width_pt: float, size: int, gap_em: float = STANZA_LINE_GAP_EM) -> float:
    """Estimated height in points of `lines`, one paragraph each, wrapped to `width_pt`."""
    rows = sum(_wrapped_line_count(line, width_pt, size, CHAR_WIDTH_EM) for line in lines)
    return rows * size * LINE_HEIGHT_EM + len(lines) * size * gap_em


def fit_font_size(lines: list[str], width_pt: float, height_pt: float, size: int, min_size: int,
                  gap_em: float = STANZA_LINE_GAP_EM) -> int:
    """Largest size <= `size` at which the lines fit the box, but never below `min_size`."""
    for candidate in range(size, min_size - 1, -1):
        if text_height(lines, width_pt, candidate, gap_em) <= height_pt:
            return candidate
    return min_size


def fit_single_line(text: str, width_pt: float, size: int, min_size: int,
                    char_em: float = BOLD_CHAR_WIDTH_EM) -> int:
    """Largest size <= `size` at which `text` (bold by default) stays on one line."""
    for candidate in range(size, min_size - 1, -1):
        if _wrapped_line_count(text, width_pt, candidate, char_em) == 1:
            return candidate
    return min_size


# --- building --------------------------------------------------------------

def _clamp(value: int, low: int, high: int) -> int:
    return max(low, min(high, value))


def _pt(length) -> float:
    return Emu(length).pt


def _remove_all_slides(prs):
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def _pick_layout(prs):
    """Prefer a layout named "Blank", otherwise the one with the fewest placeholders."""
    layouts = list(prs.slide_layouts)
    for layout in layouts:
        if layout.name.strip().lower() == "blank":
            return layout
    return min(layouts, key=lambda layout: len(layout.placeholders))


def _new_slide(prs, layout):
    slide = prs.slides.add_slide(layout)
    # Text goes in our own boxes; drop any empty "Click to add..." placeholders.
    for placeholder in list(slide.placeholders):
        placeholder._element.getparent().remove(placeholder._element)
    return slide


def _text_box(slide, left, top, width, height, inset=TEXT_INSET, wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.margin_left = inset
    frame.margin_right = inset
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.word_wrap = wrap
    return frame


def _write(paragraph, text, font_family, size, color, bold=False, align=PP_ALIGN.LEFT, space_after=None):
    run = paragraph.add_run()
    run.text = text
    run.font.name = font_family
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    paragraph.alignment = align
    if space_after is not None:
        paragraph.space_after = Pt(space_after)
    return run


def _add_title(slide, text, style, width):
    frame = _text_box(slide, SIDE_MARGIN, TITLE_TOP, width, TITLE_HEIGHT)
    size = fit_single_line(text, _pt(width - 2 * TEXT_INSET), style.title_size,
                           min(MIN_FITTED_TITLE_SIZE, style.title_size))
    _write(frame.paragraphs[0], text, style.font_family, size, BLACK, bold=True)


def _add_song_slide(prs, layout, song_title, lines, number, count, style, summary):
    slide = _new_slide(prs, layout)
    width = prs.slide_width - SIDE_MARGIN - COUNTER_RIGHT
    _add_title(slide, song_title, style, width)

    counter = _text_box(slide, prs.slide_width - COUNTER_RIGHT, TITLE_TOP, COUNTER_WIDTH, TITLE_HEIGHT,
                        inset=Inches(0.1), wrap=False)
    _write(counter.paragraphs[0], f"{number}/{count}", style.font_family, COUNTER_SIZE,
           COUNTER_BROWN, bold=True, align=PP_ALIGN.RIGHT)

    if lines:
        box_width = prs.slide_width - 2 * SIDE_MARGIN
        box_height = prs.slide_height - CONTENT_TOP - BOTTOM_MARGIN
        frame = _text_box(slide, SIDE_MARGIN, CONTENT_TOP, box_width, box_height)
        text_width = _pt(box_width - 2 * TEXT_INSET)
        text_room = _pt(box_height - Inches(0.1))  # minus the frame's top/bottom insets
        size = fit_font_size(lines, text_width, text_room, style.content_size,
                             min(MIN_FITTED_CONTENT_SIZE, style.content_size))
        if text_height(lines, text_width, size) > text_room:
            summary.warnings.append(f'"{song_title}" slide {number}: stanza is too long to fit; '
                                    "split it with a blank line")
        gap = round(size * STANZA_LINE_GAP_EM)
        for i, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            _write(paragraph, line, style.font_family, size, BLACK, space_after=gap)
    return slide


def _link_to_slide(run, source_slide, target_slide):
    """Make `run` an internal hyperlink that jumps to `target_slide` when clicked."""
    rId = source_slide.part.relate_to(target_slide.part, RT.SLIDE)
    click = run._r.get_or_add_rPr().add_hlinkClick(rId)
    click.set("action", "ppaction://hlinksldjump")


def _add_toc_slides(prs, layout, entries, style):
    """entries: (title, first slide) per song. Returns the number of TOC slides added."""
    pages = math.ceil(len(entries) / TOC_SONGS_PER_SLIDE)
    column_width = (prs.slide_width - 2 * SIDE_MARGIN - TOC_COLUMN_GAP * (TOC_COLUMNS - 1)) // TOC_COLUMNS
    column_height = prs.slide_height - TOC_TOP - TOC_BOTTOM_MARGIN
    text_width = _pt(column_width - TEXT_INSET)  # TOC columns have no right inset
    labels = [f"{n:2d}. {title}" for n, (title, _) in enumerate(entries, start=1)]
    columns = [labels[i:i + TOC_SONGS_PER_COLUMN] for i in range(0, len(labels), TOC_SONGS_PER_COLUMN)]
    # One entry size for every TOC page: the largest that lets each column fit.
    size = min(fit_font_size(column, text_width, _pt(column_height),
                             TOC_ENTRY_SIZE, TOC_MIN_ENTRY_SIZE, gap_em=TOC_LINE_GAP_EM)
               for column in columns)

    for page in range(pages):
        slide = _new_slide(prs, layout)
        heading = "Table of Contents" + (f" ({page + 1}/{pages})" if pages > 1 else "")
        _add_title(slide, heading, style, prs.slide_width - 2 * SIDE_MARGIN)

        for column in range(TOC_COLUMNS):
            first = page * TOC_SONGS_PER_SLIDE + column * TOC_SONGS_PER_COLUMN
            last = first + TOC_SONGS_PER_COLUMN
            column_entries = list(zip(labels[first:last], entries[first:last]))
            if not column_entries:
                break
            left = SIDE_MARGIN + column * (column_width + TOC_COLUMN_GAP)
            frame = _text_box(slide, left, TOC_TOP, column_width, column_height)
            frame.margin_right = 0
            for i, (label, (_, target)) in enumerate(column_entries):
                paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                entry_size = fit_single_line(label, text_width, size,
                                             min(TOC_MIN_SINGLE_LINE_SIZE, size), TOC_CHAR_WIDTH_EM)
                run = _write(paragraph, label, style.font_family, entry_size, LINK_BLUE,
                             space_after=round(size * TOC_LINE_GAP_EM))
                _link_to_slide(run, slide, target)
    return pages


def _move_last_slides_to_front(prs, count):
    slide_ids = prs.slides._sldIdLst
    moved = list(slide_ids)[-count:]
    for slide_id in moved:
        slide_ids.remove(slide_id)
    for position, slide_id in enumerate(moved):
        slide_ids.insert(position, slide_id)


def build_presentation(songs, template=None, *, toc=False, style=None):
    """Build the deck. `template` is a .pptx path, file-like object, or None.

    The template's own slides are removed; only its masters, layouts and size are used.
    Returns (Presentation, Summary).
    """
    style = style or Style()
    prs = Presentation(template)
    _remove_all_slides(prs)
    layout = _pick_layout(prs)
    summary = Summary(songs=len(songs))

    entries = []
    for song in songs:
        stanzas = song.stanzas or [[]]
        if not song.stanzas:
            summary.warnings.append(f'"{song.title}" has no lyrics; it gets a title-only slide')
        for number, lines in enumerate(stanzas, start=1):
            slide = _add_song_slide(prs, layout, song.title, lines, number, len(stanzas), style, summary)
            summary.song_slides += 1
            if number == 1:
                entries.append((song.title, slide))

    if toc and entries:
        summary.toc_slides = _add_toc_slides(prs, layout, entries, style)
        _move_last_slides_to_front(prs, summary.toc_slides)
    return prs, summary


def generate_pptx(text, template_bytes=None, *, toc=False, style=None):
    """Song text (+ optional template bytes) -> (.pptx bytes, Summary). Used by the browser app."""
    songs = parse_songs(text)
    if not songs:
        raise ValueError('No songs found. Start each song with a title line such as "# Amazing Grace".')
    template = io.BytesIO(template_bytes) if template_bytes else None
    prs, summary = build_presentation(songs, template, toc=toc, style=style)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), summary


# --- template preview (browser) ------------------------------------------------
#
# The browser preview has no PowerPoint renderer, so we draw the template's own
# artwork (background, master and layout shapes) as SVG and lay the slide text on
# top. This covers what church templates use: fills, pictures, lines, rectangles,
# ellipses, custom paths, groups and simple text. Placeholders are skipped because
# the generator never fills them.

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
EMU_PER_PT = 12700
WEB_IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/svg+xml", "image/webp", "image/bmp"}
# Free fonts with the same character widths, used when the real font is not installed.
METRIC_TWINS = {"Calibri": "Carlito", "Arial": "Arimo", "Times New Roman": "Tinos", "Georgia": "Gelasio"}


def _local(el) -> str:
    return el.tag.rsplit("}", 1)[-1] if isinstance(el.tag, str) else ""


def _font_stack(face: str) -> str:
    twin = METRIC_TWINS.get(face)
    return ", ".join([f"'{face}'"] + ([f"'{twin}'"] if twin else []) + ["sans-serif"])


class _Colors:
    """Resolves DrawingML colours (srgb, scheme, sys, preset + lum/tint/shade/alpha)."""

    PRESET = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000", "blue": "0000FF"}

    def __init__(self, master):
        from lxml import etree
        self.scheme, self.minor_font = {}, "Calibri"
        theme = master.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme.blob)
        clr_scheme = root.find(f".//{_A}clrScheme")
        if clr_scheme is not None:
            for slot in clr_scheme:
                value = slot[0] if len(slot) else None
                if value is not None:
                    self.scheme[_local(slot)] = value.get("val") if _local(value) == "srgbClr" else value.get("lastClr")
        minor = root.find(f".//{_A}minorFont/{_A}latin")
        if minor is not None and minor.get("typeface"):
            self.minor_font = minor.get("typeface")
        clr_map = master._element.find(f"{_P}clrMap")
        self.map = dict(clr_map.attrib) if clr_map is not None else {}

    def font(self, face):
        if not face or face.startswith("+mn") or face.startswith("+mj"):
            return self.minor_font
        return face

    def of(self, parent):
        """Colour of the first colour child of `parent`: ("#RRGGBB", opacity) or None."""
        if parent is None:
            return None
        for el in parent:
            kind = _local(el)
            if kind == "srgbClr":
                hex_ = el.get("val")
            elif kind == "schemeClr":
                name = el.get("val")
                hex_ = self.scheme.get(self.map.get(name, name)) or self.scheme.get(name)
            elif kind == "sysClr":
                hex_ = el.get("lastClr") or ("000000" if el.get("val") == "windowText" else "FFFFFF")
            elif kind == "prstClr":
                hex_ = self.PRESET.get(el.get("val"), "000000")
            else:
                continue
            return self._modify(hex_ or "000000", el)
        return None

    @staticmethod
    def _modify(hex_, el):
        import colorsys
        r, g, b = (int(hex_[i:i + 2], 16) / 255 for i in (0, 2, 4))
        alpha = 1.0
        for mod in el:
            kind, val = _local(mod), int(mod.get("val", "100000")) / 100000
            if kind in ("lumMod", "lumOff"):
                h, l, s = colorsys.rgb_to_hls(r, g, b)
                l = min(1, max(0, l * val if kind == "lumMod" else l + val))
                r, g, b = colorsys.hls_to_rgb(h, l, s)
            elif kind == "shade":
                r, g, b = r * val, g * val, b * val
            elif kind == "tint":
                r, g, b = (c + (1 - c) * (1 - val) for c in (r, g, b))
            elif kind == "alpha":
                alpha = val
        return "#" + "".join(f"{round(c * 255):02X}" for c in (r, g, b)), alpha


def _paint(attr, colour):
    if colour is None:
        return f'{attr}="none"'
    hex_, alpha = colour
    return f'{attr}="{hex_}"' + (f' {attr}-opacity="{alpha:g}"' if alpha < 1 else "")


def _fill_of(sp_pr, style, colors):
    if sp_pr is not None:
        for el in sp_pr:
            kind = _local(el)
            if kind == "noFill":
                return None
            if kind == "solidFill":
                return colors.of(el)
            if kind == "gradFill":
                stop = el.find(f"{_A}gsLst/{_A}gs")
                return colors.of(stop)
    ref = style.find(f"{_A}fillRef") if style is not None else None
    if ref is not None and ref.get("idx", "0") != "0":
        return colors.of(ref)
    return None


def _line_of(sp_pr, style, colors):
    ln = sp_pr.find(f"{_A}ln") if sp_pr is not None else None
    width = int(ln.get("w", "9525")) if ln is not None else 9525
    if ln is not None:
        if ln.find(f"{_A}noFill") is not None:
            return None, 0
        solid = ln.find(f"{_A}solidFill")
        if solid is not None:
            return colors.of(solid), width
    ref = style.find(f"{_A}lnRef") if style is not None else None
    if ref is not None and ref.get("idx", "0") != "0":
        return colors.of(ref), width
    return None, 0


def _xfrm(sp_pr):
    xfrm = sp_pr.find(f"{_A}xfrm") if sp_pr is not None else None
    if xfrm is None or xfrm.find(f"{_A}off") is None:
        return None
    off, ext = xfrm.find(f"{_A}off"), xfrm.find(f"{_A}ext")
    return {
        "x": int(off.get("x")), "y": int(off.get("y")),
        "cx": int(ext.get("cx")), "cy": int(ext.get("cy")),
        "rot": int(xfrm.get("rot", "0")) / 60000,
        "flipH": xfrm.get("flipH") in ("1", "true"), "flipV": xfrm.get("flipV") in ("1", "true"),
        "el": xfrm,
    }


def _transform(box):
    cx, cy = box["x"] + box["cx"] / 2, box["y"] + box["cy"] / 2
    parts = []
    if box["rot"]:
        parts.append(f"rotate({box['rot']:g} {cx:.0f} {cy:.0f})")
    if box["flipH"] or box["flipV"]:
        sx, sy = (-1 if box["flipH"] else 1), (-1 if box["flipV"] else 1)
        parts.append(f"translate({cx:.0f} {cy:.0f}) scale({sx} {sy}) translate({-cx:.0f} {-cy:.0f})")
    return f' transform="{" ".join(parts)}"' if parts else ""


def _path_data(cust_geom, box):
    out = []
    for path in cust_geom.iter(f"{_A}path"):
        w = int(path.get("w", box["cx"]) or box["cx"]) or 1
        h = int(path.get("h", box["cy"]) or box["cy"]) or 1
        sx, sy = box["cx"] / w, box["cy"] / h
        point = lambda pt: f"{box['x'] + int(pt.get('x')) * sx:.0f} {box['y'] + int(pt.get('y')) * sy:.0f}"
        d = []
        for cmd in path:
            kind, pts = _local(cmd), cmd.findall(f"{_A}pt")
            if kind == "moveTo":
                d.append("M " + point(pts[0]))
            elif kind == "lnTo":
                d.append("L " + point(pts[0]))
            elif kind == "cubicBezTo":
                d.append("C " + ", ".join(point(p) for p in pts))
            elif kind == "quadBezTo":
                d.append("Q " + ", ".join(point(p) for p in pts))
            elif kind == "close":
                d.append("Z")
        out.append((" ".join(d), path.get("fill") != "none", path.get("stroke") not in ("0", "false")))
    return out


def _text_svg(sp, box, colors, master_default_size):
    from xml.sax.saxutils import escape
    tx_body = sp.find(f"{_P}txBody")
    if tx_body is None:
        return ""
    body = tx_body.find(f"{_A}bodyPr")
    lst = tx_body.find(f"{_A}lstStyle/{_A}lvl1pPr")
    lst_rpr = lst.find(f"{_A}defRPr") if lst is not None else None
    ins = lambda name, default: int(body.get(name, default)) if body is not None else default
    left, top, right, bottom = ins("lIns", 91440), ins("tIns", 45720), ins("rIns", 91440), ins("bIns", 45720)

    lines = []
    for p in tx_body.findall(f"{_A}p"):
        runs = [r for r in p if _local(r) in ("r", "fld")]
        text = "".join(r.findtext(f"{_A}t") or "" for r in runs)
        if not text.strip():
            continue
        rpr = runs[0].find(f"{_A}rPr")
        pick = lambda attr, default=None: (rpr.get(attr) if rpr is not None and rpr.get(attr) else
                                           lst_rpr.get(attr) if lst_rpr is not None and lst_rpr.get(attr) else default)
        size = int(pick("sz", master_default_size)) / 100
        bold = pick("b", "0") in ("1", "true")
        colour = (colors.of(rpr.find(f"{_A}solidFill")) if rpr is not None and rpr.find(f"{_A}solidFill") is not None
                  else colors.of(lst_rpr.find(f"{_A}solidFill")) if lst_rpr is not None and lst_rpr.find(f"{_A}solidFill") is not None
                  else colors.of(_default_text_colour()))
        latin = (rpr.find(f"{_A}latin") if rpr is not None else None)
        if latin is None and lst_rpr is not None:
            latin = lst_rpr.find(f"{_A}latin")
        face = colors.font(latin.get("typeface") if latin is not None else None)
        ppr = p.find(f"{_A}pPr")
        align = (ppr.get("algn") if ppr is not None and ppr.get("algn") else
                 lst.get("algn") if lst is not None and lst.get("algn") else "l")
        lines.append((text, size, bold, colour, face, align))
    if not lines:
        return ""

    line_heights = [size * EMU_PER_PT * 1.2 for _, size, *_ in lines]
    anchor = body.get("anchor", "t") if body is not None else "t"
    y = box["y"] + top
    room = box["cy"] - top - bottom
    if anchor == "ctr":
        y += (room - sum(line_heights)) / 2
    elif anchor == "b":
        y += room - sum(line_heights)
    out = []
    for (text, size, bold, colour, face, align), height in zip(lines, line_heights):
        x, text_anchor = box["x"] + left, "start"
        if align == "ctr":
            x, text_anchor = box["x"] + box["cx"] / 2, "middle"
        elif align == "r":
            x, text_anchor = box["x"] + box["cx"] - right, "end"
        baseline = y + size * EMU_PER_PT * 0.95
        # sized in points and scaled up: browsers cap font-size, and EMU sizes run past it
        out.append(f'<text transform="translate({x:.0f} {baseline:.0f}) scale({EMU_PER_PT})" font-size="{size:g}" '
                   f'font-family="{_font_stack(face)}" font-weight="{700 if bold else 400}" '
                   f'text-anchor="{text_anchor}" {_paint("fill", colour)}>{escape(text)}</text>')
        y += height
    return "".join(out)


def _default_text_colour():
    """A fill holding <a:schemeClr val="tx1"/>: the theme's default text colour."""
    from lxml import etree
    holder = etree.Element(f"{_A}solidFill")
    etree.SubElement(holder, f"{_A}schemeClr", val="tx1")
    return holder


def _shapes_svg(sp_tree, part, colors, master_default_size):
    import base64
    out = []
    for el in sp_tree:
        kind = _local(el)
        if kind not in ("sp", "pic", "cxnSp", "grpSp"):
            continue
        if el.find(f".//{_P}nvPr/{_P}ph") is not None and kind != "grpSp":
            continue  # placeholder: the generator never fills these
        sp_pr = el.find(f"{_P}spPr") if kind != "grpSp" else el.find(f"{_P}grpSpPr")
        box = _xfrm(sp_pr)
        if box is None:
            continue
        if kind == "grpSp":
            ch_off, ch_ext = box["el"].find(f"{_A}chOff"), box["el"].find(f"{_A}chExt")
            sx = box["cx"] / (int(ch_ext.get("cx")) or 1) if ch_ext is not None else 1
            sy = box["cy"] / (int(ch_ext.get("cy")) or 1) if ch_ext is not None else 1
            ox, oy = (int(ch_off.get("x")), int(ch_off.get("y"))) if ch_off is not None else (0, 0)
            inner = _shapes_svg(el, part, colors, master_default_size)
            out.append(f'<g{_transform(box)}><g transform="translate({box["x"]:.0f} {box["y"]:.0f}) '
                       f'scale({sx:g} {sy:g}) translate({-ox:.0f} {-oy:.0f})">{inner}</g></g>')
            continue
        if kind == "pic":
            blip = el.find(f".//{_A}blip")
            rid = blip.get(f"{_R}embed") if blip is not None else None
            if not rid:
                continue
            image = part.related_part(rid)
            if image.content_type not in WEB_IMAGE_TYPES:
                continue
            data = base64.b64encode(image.blob).decode("ascii")
            out.append(f'<image x="{box["x"]}" y="{box["y"]}" width="{box["cx"]}" height="{box["cy"]}" '
                       f'preserveAspectRatio="none" href="data:{image.content_type};base64,{data}"'
                       f'{_transform(box)}/>')
            continue

        style = el.find(f"{_P}style")
        fill = _fill_of(sp_pr, style, colors)
        stroke, width = _line_of(sp_pr, style, colors)
        no_stroke = 'stroke="none"'
        stroke_attrs = f'{_paint("stroke", stroke)} stroke-width="{width}"' if stroke else no_stroke
        prst = sp_pr.find(f"{_A}prstGeom")
        cust = sp_pr.find(f"{_A}custGeom")
        geometry = prst.get("prst") if prst is not None else None
        x, y, cx, cy = box["x"], box["y"], box["cx"], box["cy"]
        if kind == "cxnSp" or geometry in ("line", "straightConnector1"):
            x1, x2 = (x + cx, x) if box["flipH"] else (x, x + cx)
            y1, y2 = (y + cy, y) if box["flipV"] else (y, y + cy)
            if stroke:
                rotate = f' transform="rotate({box["rot"]:g} {x + cx / 2:.0f} {y + cy / 2:.0f})"' if box["rot"] else ""
                out.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" {stroke_attrs}{rotate}/>')
        elif cust is not None:
            for d, can_fill, can_stroke in _path_data(cust, box):
                out.append(f'<path d="{d}" {_paint("fill", fill if can_fill else None)} '
                           f'{stroke_attrs if can_stroke else no_stroke}{_transform(box)}/>')
        elif fill or stroke:
            if geometry == "ellipse":
                out.append(f'<ellipse cx="{x + cx / 2:.0f}" cy="{y + cy / 2:.0f}" rx="{cx / 2:.0f}" ry="{cy / 2:.0f}" '
                           f'{_paint("fill", fill)} {stroke_attrs}{_transform(box)}/>')
            else:
                radius = min(cx, cy) * 0.1667 if geometry == "roundRect" else 0
                out.append(f'<rect x="{x}" y="{y}" width="{cx}" height="{cy}" rx="{radius:.0f}" '
                           f'{_paint("fill", fill)} {stroke_attrs}{_transform(box)}/>')
        out.append(_text_svg(el, box, colors, master_default_size))
    return "".join(out)


def _background_svg(holder, part, colors, width, height):
    import base64
    bg = holder._element.find(f"{_P}cSld/{_P}bg") if holder is not None else None
    if bg is None:
        return None
    bg_pr, bg_ref = bg.find(f"{_P}bgPr"), bg.find(f"{_P}bgRef")
    if bg_pr is not None:
        blip = bg_pr.find(f".//{_A}blip")
        if blip is not None and blip.get(f"{_R}embed"):
            image = part.related_part(blip.get(f"{_R}embed"))
            if image.content_type in WEB_IMAGE_TYPES:
                data = base64.b64encode(image.blob).decode("ascii")
                return (f'<image x="0" y="0" width="{width}" height="{height}" preserveAspectRatio="none" '
                        f'href="data:{image.content_type};base64,{data}"/>')
        colour = _fill_of(bg_pr, None, colors)
    else:
        colour = colors.of(bg_ref)
    return f'<rect x="0" y="0" width="{width}" height="{height}" {_paint("fill", colour)}/>' if colour else None


def template_preview(template=None) -> dict:
    """Draw the template's artwork as SVG for the browser preview.

    `template` is .pptx bytes, a path, a file-like object or None. Returns
    {"width_in", "height_in", "svg"}; the SVG's user units are EMU.
    """
    if isinstance(template, (bytes, bytearray)):
        template = io.BytesIO(template)
    prs = Presentation(template)
    layout = _pick_layout(prs)
    master = layout.slide_master
    colors = _Colors(master)
    width, height = prs.slide_width, prs.slide_height

    other = master._element.find(f"{_P}txStyles/{_P}otherStyle/{_A}lvl1pPr/{_A}defRPr")
    default_size = int(other.get("sz")) if other is not None and other.get("sz") else 1800

    layers = [_background_svg(layout, layout.part, colors, width, height)
              or _background_svg(master, master.part, colors, width, height)
              or f'<rect x="0" y="0" width="{width}" height="{height}" fill="#FFFFFF"/>']
    if layout._element.get("showMasterSp") not in ("0", "false"):
        layers.append(_shapes_svg(master._element.find(f"{_P}cSld/{_P}spTree"), master.part, colors, default_size))
    layers.append(_shapes_svg(layout._element.find(f"{_P}cSld/{_P}spTree"), layout.part, colors, default_size))
    svg = (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" '
           f'preserveAspectRatio="none">{"".join(layers)}</svg>')
    # PowerPoint and LibreOffice colour hyperlinks with the theme's hlink colour, not the run's.
    link = "#" + colors.scheme.get("hlink", "0000FF")
    return {"width_in": round(Emu(width).inches, 4), "height_in": round(Emu(height).inches, 4),
            "svg": svg, "link": link}
